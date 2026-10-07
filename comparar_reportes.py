"""
Checks that the refactor didn't change the output.
Run the OLD app and the NEW app on the same PDFs + same blank Excel, download both
reports, then:
    python comparar_reportes.py Reporte_viejo.xlsx Reporte_nuevo.xlsx
"""
import sys
import openpyxl

viejo = openpyxl.load_workbook(sys.argv[1])
nuevo = openpyxl.load_workbook(sys.argv[2])
diferencias = 0

if viejo.sheetnames != nuevo.sheetnames:
    print(f"Hojas distintas: {viejo.sheetnames} vs {nuevo.sheetnames}")
    diferencias += 1

for nombre in viejo.sheetnames:
    if nombre not in nuevo.sheetnames:
        continue
    a, b = viejo[nombre], nuevo[nombre]
    for r in range(1, max(a.max_row, b.max_row) + 1):
        for c in range(1, max(a.max_column, b.max_column) + 1):
            va, vb = a.cell(r, c).value, b.cell(r, c).value
            if va != vb:
                print(f"[{nombre}] {a.cell(r, c).coordinate}: viejo={va!r}  nuevo={vb!r}")
                diferencias += 1

print("Idénticos ✅" if diferencias == 0 else f"{diferencias} diferencia(s) encontradas")
