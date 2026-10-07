"""
Writes processed facturas into MAGA's Excel template.
Usage:
    reporte = ReporteMAGA(xlsx_bytes, EXCEL_MAPPINGS)   # raises PlantillaInvalidaError
    reporte.agregar_factura(factura, m_id, m_name)       # once per factura
    xlsx_out = reporte.finalizar()                        # bytes for download
"""
import io
import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from texto import normalize_text, squish_text, safe_float
from config import UMBRAL_ALERTA_ABARROTES


class PlantillaInvalidaError(ValueError):
    """The uploaded Excel doesn't have the columns the report needs."""


def get_master_cell(ws, r_idx, c_idx):
    cell = ws.cell(row=r_idx, column=c_idx)
    if type(cell).__name__ == 'MergedCell':
        for m_range in ws.merged_cells.ranges:
            if cell.coordinate in m_range:
                return ws.cell(row=m_range.min_row, column=m_range.min_col)
    return cell


def preparar_hojas(wb):
    """Creates (or reuses) the 'Extra Detalles' and 'Items Sin Clasificar' sheets."""
    if "Extra Detalles" not in wb.sheetnames:
        ws_det = wb.create_sheet("Extra Detalles")
        ws_det.append(['Archivo PDF', 'Nombre Emisor', 'NIT Emisor', 'NIT Receptor', 'Nombre Escuela', 'Num. DTE', 'Municipio', 'Alerta % Abarrotes'])
    else:
        ws_det = wb["Extra Detalles"]
        # Excel files from older runs lack the 'Nombre Escuela' column:
        # insert it after 'NIT Receptor' so old and new rows stay aligned.
        if normalize_text(str(ws_det.cell(row=1, column=5).value or "")) != normalize_text("Nombre Escuela"):
            ws_det.insert_cols(5)
            ws_det.cell(row=1, column=5).value = "Nombre Escuela"

    if "Items Sin Clasificar" not in wb.sheetnames:
        ws_unmatched = wb.create_sheet("Items Sin Clasificar")
        ws_unmatched.append(['Descripción', 'Municipio', 'Total (Q)', 'Num. DTE'])
    else:
        ws_unmatched = wb["Items Sin Clasificar"]

    return ws_det, ws_unmatched


def mapear_columnas(ws):
    """Finds the abarrotes / agricultura / escuelas / productores columns by header text."""
    col_map = {}
    for row in ws.iter_rows(min_row=1, max_row=15):
        for cell in row:
            if type(cell).__name__ == 'MergedCell': continue
            if not cell.value: continue
            val = normalize_text(str(cell.value))

            if 'abarrotes' in val: col_map['abar'] = cell.column
            if 'agricultura' in val: col_map['agri'] = cell.column
            if 'escuela' in val or 'establecimiento' in val: col_map['escuelas'] = cell.column
            if 'proveedor' in val or 'productor' in val:
                base_col, base_row, found_total = cell.column, cell.row, False
                for r_offset in range(1, 4):
                    for c_offset in range(3):
                        sub_cell = ws.cell(row=base_row + r_offset, column=base_col + c_offset)
                        if sub_cell.value and 'total' in normalize_text(str(sub_cell.value)):
                            col_map['productores'] = sub_cell.column
                            found_total = True
                            break
                    if found_total: break
                if 'productores' not in col_map: col_map['productores'] = base_col
    return col_map


def mapear_filas(ws, excel_mappings):
    """Finds the row of each municipality in the main sheet."""
    row_map = {}
    for row_ex in ws.iter_rows(min_row=5, max_row=150):
        row_text = " ".join([str(c.value) for c in row_ex if c.value and type(c).__name__ != 'MergedCell'])
        row_squished = squish_text(row_text)
        for m_id, search_key in excel_mappings.items():
            if m_id in row_map: continue
            if squish_text(search_key) in row_squished:
                row_map[m_id] = row_ex[0].row
    return row_map


def formatear_hoja(ws):
    """Thin borders and auto column width."""
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.border = thin_border
            try: max_length = max(max_length, len(str(cell.value)))
            except: pass
        ws.column_dimensions[col_letter].width = max_length + 2


class ReporteMAGA:
    def __init__(self, xlsx_bytes, excel_mappings):
        self.wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.ws = self.wb.active
        self.ws_det, self.ws_unmatched = preparar_hojas(self.wb)

        self.col_map = mapear_columnas(self.ws)
        if 'abar' not in self.col_map or 'agri' not in self.col_map:
            raise PlantillaInvalidaError("No encontré las columnas base en el Excel.")

        self.row_map = mapear_filas(self.ws, excel_mappings)
        self.batch_totals = {}

    def agregar_factura(self, factura, m_id, m_name):
        """Adds one classified factura (after clasificacion.clasificar_factura)."""
        for linea in factura['lineas']:
            if linea['categoria'] == 'unmatched':
                self.ws_unmatched.append([linea['descripcion'], m_name, linea['total'], factura['dte']])

        totales = self.batch_totals.setdefault(
            m_id, {'abar': 0.0, 'agri': 0.0, 'emisores': set(), 'receptores': set()})
        abar_sum, agri_sum = factura['total_abar'], factura['total_agri']
        totales['abar'] += abar_sum
        totales['agri'] += agri_sum
        if factura['nit_emisor'] != "N/A": totales['emisores'].add(factura['nit_emisor'])
        if factura['nit_receptor'] != "N/A": totales['receptores'].add(factura['nit_receptor'])

        total_rec = abar_sum + agri_sum
        perc_abar = (abar_sum / total_rec) if total_rec > 0 else 0
        alert_status = "⚠️ ALERTA: >30%" if perc_abar > UMBRAL_ALERTA_ABARROTES else "OK"

        self.ws_det.append([factura['archivo'], factura['nombre_emisor'], factura['nit_emisor'],
                            factura['nit_receptor'], factura['nombre_escuela'], factura['dte'],
                            m_name, alert_status])

    @property
    def unmatched_count(self):
        # Includes rows left over from earlier runs, as in the original tool
        return self.ws_unmatched.max_row - 1 if self.ws_unmatched.max_row > 1 else 0

    def finalizar(self):
        """Adds batch totals to the main sheet, formats the extra sheets, returns xlsx bytes."""
        for target_m_id, r_idx in self.row_map.items():
            data = self.batch_totals.get(target_m_id)
            if not data: continue

            if 'abar' in self.col_map and data['abar'] > 0:
                target_cell = get_master_cell(self.ws, r_idx, self.col_map['abar'])
                target_cell.value = safe_float(target_cell.value) + data['abar']

            if 'agri' in self.col_map and data['agri'] > 0:
                target_cell = get_master_cell(self.ws, r_idx, self.col_map['agri'])
                target_cell.value = safe_float(target_cell.value) + data['agri']

            if 'escuelas' in self.col_map and len(data['receptores']) > 0:
                target_cell = get_master_cell(self.ws, r_idx, self.col_map['escuelas'])
                target_cell.value = int(safe_float(target_cell.value)) + len(data['receptores'])

            if 'productores' in self.col_map and len(data['emisores']) > 0:
                target_cell = get_master_cell(self.ws, r_idx, self.col_map['productores'])
                target_cell.value = int(safe_float(target_cell.value)) + len(data['emisores'])

        formatear_hoja(self.ws_det)
        formatear_hoja(self.ws_unmatched)

        output = io.BytesIO()
        self.wb.save(output)
        return output.getvalue()
